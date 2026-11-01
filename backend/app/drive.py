"""Google Drive read-only OAuth and on-demand source retrieval."""

import hashlib
import io
import logging
import os
import secrets
import time
from urllib.parse import urlencode

import httpx
import pypdfium2 as pdfium
from docx import Document as WordDocument
from fastapi import APIRouter, HTTPException, Request
from fastapi.responses import RedirectResponse
from pydantic import BaseModel, Field

from .db import DriveOAuthState, ExternalConnection, Session, SourceHealthCheck, User, Workspace
from .gmail import actor, allowed_workspace, crypt

router = APIRouter(prefix='/api')
log = logging.getLogger(__name__)
SCOPE = 'https://www.googleapis.com/auth/drive.readonly'
TOKEN_URL = 'https://oauth2.googleapis.com/token'
DRIVE_API = 'https://www.googleapis.com/drive/v3/files'
CLIENT_ID = os.getenv('GOOGLE_CLIENT_ID', '')
CLIENT_SECRET = os.getenv('GOOGLE_CLIENT_SECRET', '')
REDIRECT = os.getenv('GOOGLE_DRIVE_REDIRECT_URI', 'http://localhost:3000/api/drive/callback')


def configured():
    if not CLIENT_ID or not CLIENT_SECRET or not REDIRECT:
        raise HTTPException(503, 'Configure Google OAuth credentials for Drive.')
    crypt()


def request_google(url, access_token, params=None, *, json_response=True, binary=False):
    try:
        response = httpx.get(
            url, params=params, headers={'Authorization': 'Bearer ' + access_token}, timeout=20
        )
        response.raise_for_status()
        if binary:
            if len(response.content) > 12 * 1024 * 1024:
                raise HTTPException(413, 'Drive file exceeds the 12 MB evidence limit.')
            return response.content
        return response.json() if json_response else response.text
    except httpx.HTTPStatusError as exc:
        status = exc.response.status_code
        reason = ''
        try:
            payload = exc.response.json().get('error', {})
            reason = next(
                (
                    item.get('reason', '')
                    for item in payload.get('errors', [])
                    if item.get('reason')
                ),
                '',
            )
        except (ValueError, TypeError, AttributeError):
            pass
        log.warning('Google Drive request rejected: HTTP %s, reason %s', status, reason)
        if status == 401:
            raise HTTPException(
                401, 'Drive authorization expired or was revoked. Reconnect Google Drive.'
            ) from exc
        if status == 403:
            if reason in ('accessNotConfigured', 'SERVICE_DISABLED'):
                raise HTTPException(
                    503, 'Enable Google Drive API in the OAuth client Google Cloud project.'
                ) from exc
            if reason in ('insufficientPermissions', 'ACCESS_TOKEN_SCOPE_INSUFFICIENT'):
                raise HTTPException(
                    403,
                    'Drive read permission was not granted. Disconnect and reconnect Google Drive.',
                ) from exc
            raise HTTPException(
                403,
                'Google Drive denied access. Confirm its API is enabled and inspect the backend log reason.',
            ) from exc
        if exc.response.status_code == 429:
            raise HTTPException(429, 'Google Drive quota exceeded. Retry later.') from exc
        raise HTTPException(502, 'Google Drive returned an error.') from exc
    except (httpx.RequestError, ValueError) as exc:
        raise HTTPException(502, 'Could not read Google Drive.') from exc


def google_token(payload):
    try:
        response = httpx.post(TOKEN_URL, data=payload, timeout=20)
        response.raise_for_status()
        return response.json()
    except (httpx.HTTPError, ValueError) as exc:
        raise HTTPException(
            502, 'Google authorization or refresh failed. Reconnect Drive.'
        ) from exc


@router.get('/drive/connect')
def connect(workspaceId: str, request: Request):
    configured()
    with Session() as db:
        who = actor(request, db)
        allowed_workspace(db, who, workspaceId)
        state = secrets.token_urlsafe(32)
        db.add(
            DriveOAuthState(
                value=hashlib.sha256(state.encode()).hexdigest(),
                user_id=who.id,
                workspace_id=workspaceId,
                expires=time.time() + 600,
            )
        )
        db.commit()
    query = {
        'client_id': CLIENT_ID,
        'redirect_uri': REDIRECT,
        'response_type': 'code',
        'scope': SCOPE,
        'state': state,
        'access_type': 'offline',
        'prompt': 'consent',
    }
    return {'authorization_url': 'https://accounts.google.com/o/oauth2/v2/auth?' + urlencode(query)}


@router.get('/drive/callback')
def callback(state: str = '', code: str = ''):
    configured()
    if not state or not code:
        raise HTTPException(400, 'Drive authorization was cancelled.')
    with Session() as db:
        item = db.get(DriveOAuthState, hashlib.sha256(state.encode()).hexdigest())
        if not item or item.expires < time.time():
            raise HTTPException(400, 'Drive authorization expired. Start again.')
        workspaceId = item.workspace_id
        who = db.get(User, item.user_id)
        if not who or not db.query(Workspace).filter_by(id=workspaceId, owner_id=who.id).first():
            raise HTTPException(400, 'Workspace no longer exists.')
        db.delete(item)
        db.commit()
        tokens = google_token(
            {
                'code': code,
                'client_id': CLIENT_ID,
                'client_secret': CLIENT_SECRET,
                'redirect_uri': REDIRECT,
                'grant_type': 'authorization_code',
            }
        )
        refresh = tokens.get('refresh_token')
        if not refresh or SCOPE not in tokens.get('scope', '').split():
            raise HTTPException(
                400, 'Drive read-only access was not granted. Try connecting again.'
            )
        access = tokens.get('access_token')
        if not access:
            raise HTTPException(502, 'Google did not return an access token.')
        identity = request_google(
            'https://www.googleapis.com/drive/v3/about',
            access,
            {'fields': 'user(displayName,emailAddress)'},
        )
        account = (identity.get('user') or {}).get('emailAddress') or 'Google Drive'
        row = db.get(ExternalConnection, (workspaceId, 'drive'))
        encrypted = crypt().encrypt(refresh.encode()).decode()
        if row:
            row.account, row.secret = account, encrypted
        else:
            db.add(
                ExternalConnection(
                    workspace_id=workspaceId, provider='drive', account=account, secret=encrypted
                )
            )
        health = db.get(SourceHealthCheck, (workspaceId, 'drive'))
        if health:
            db.delete(health)
        db.commit()
    return RedirectResponse('/?drive=connected', status_code=303)


def access_token(db, who, workspaceId):
    allowed_workspace(db, who, workspaceId)
    row = db.get(ExternalConnection, (workspaceId, 'drive'))
    if not row:
        raise HTTPException(409, 'Connect Google Drive first.')
    try:
        refresh = crypt().decrypt(row.secret.encode()).decode()
    except Exception as exc:
        raise HTTPException(503, 'Drive credentials cannot be decrypted. Reconnect Drive.') from exc
    tokens = google_token(
        {
            'client_id': CLIENT_ID,
            'client_secret': CLIENT_SECRET,
            'refresh_token': refresh,
            'grant_type': 'refresh_token',
        }
    )
    return tokens['access_token']


def drive_search(token, query):
    term = query.strip().replace('\\', '\\\\').replace("'", "\\'")
    q = 'trashed = false'
    if term:
        q += " and (name contains '" + term[:100] + "' or fullText contains '" + term[:100] + "')"
    data = request_google(
        DRIVE_API,
        token,
        {
            'q': q,
            'page_size': 20,
            'fields': 'nextPageToken,files(id,name,mimeType,webViewLink,description,modifiedTime)',
            'orderBy': 'modifiedTime desc',
        },
    )
    return [
        {
            'provider': 'drive',
            'id': item['id'],
            'title': item.get('name', 'Untitled'),
            'excerpt': item.get('description') or item.get('name', ''),
            'url': item.get('webViewLink')
            or 'https://drive.google.com/file/d/' + item['id'] + '/view',
            'mimeType': item.get('mimeType', ''),
        }
        for item in data.get('files', [])
        if item.get('id')
    ]


def extract_drive_binary(body, mime):
    """Small, on-demand extraction without adding files to the indexed library."""
    if mime == 'application/pdf':
        pdf = pdfium.PdfDocument(body)
        parts = []
        try:
            for i in range(min(len(pdf), 20)):
                page = pdf[i]
                try:
                    textpage = page.get_textpage()
                    try:
                        parts.append(textpage.get_text_bounded())
                    finally:
                        textpage.close()
                finally:
                    page.close()
        finally:
            pdf.close()
        return '\n'.join(parts)[:6000]
    if mime == 'application/vnd.openxmlformats-officedocument.wordprocessingml.document':
        doc = WordDocument(io.BytesIO(body))
        return '\n'.join(p.text for p in doc.paragraphs if p.text.strip())[:6000]
    return ''


def drive_evidence(token, items):
    results = []
    for item in items[:5]:
        mime = item.get('mimeType', '')
        try:
            if mime == 'application/vnd.google-apps.document':
                body = request_google(
                    DRIVE_API + '/' + item['id'] + '/export',
                    token,
                    {'mimeType': 'text/plain'},
                    json_response=False,
                )
            elif mime == 'application/vnd.google-apps.spreadsheet':
                body = request_google(
                    DRIVE_API + '/' + item['id'] + '/export',
                    token,
                    {'mimeType': 'text/csv'},
                    json_response=False,
                )
            elif mime.startswith('text/') or mime in ('application/json', 'application/xml'):
                body = request_google(
                    DRIVE_API + '/' + item['id'], token, {'alt': 'media'}, json_response=False
                )
            elif mime in (
                'application/pdf',
                'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
            ):
                binary = request_google(
                    DRIVE_API + '/' + item['id'], token, {'alt': 'media'}, binary=True
                )
                body = extract_drive_binary(binary, mime)
            else:
                continue
            if body.strip():
                results.append({**item, 'excerpt': body[:1500]})
        except (HTTPException, ValueError, pdfium.PdfiumError):
            log.warning('Skipping inaccessible or unreadable Drive file %s', item['id'])
    return results


class DriveSearch(BaseModel):
    query: str = Field(default='', max_length=180)


@router.get('/workspaces/{workspaceId}/drive/status')
def status(workspaceId: str, request: Request):
    with Session() as db:
        who = actor(request, db)
        allowed_workspace(db, who, workspaceId)
        row = db.get(ExternalConnection, (workspaceId, 'drive'))
        return {
            'connected': bool(row),
            'account': row.account if row else None,
            'configured': bool(CLIENT_ID and CLIENT_SECRET),
        }


@router.post('/workspaces/{workspaceId}/drive/search')
def search(workspaceId: str, body: DriveSearch, request: Request):
    with Session() as db:
        who = actor(request, db)
        token = access_token(db, who, workspaceId)
    return {'results': drive_search(token, body.query)}


@router.delete('/workspaces/{workspaceId}/drive')
def disconnect(workspaceId: str, request: Request):
    with Session() as db:
        who = actor(request, db)
        allowed_workspace(db, who, workspaceId)
        db.query(ExternalConnection).filter_by(workspace_id=workspaceId, provider='drive').delete()
        health = db.get(SourceHealthCheck, (workspaceId, 'drive'))
        if health:
            db.delete(health)
        db.commit()
    return {'ok': True}

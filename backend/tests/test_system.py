import time
from pathlib import Path
from types import SimpleNamespace

import pytest
from fastapi.testclient import TestClient

from app import ai, config
from app.db import Chunk, Document, Session
from app.ingest import split_text
from app.main import app
from app.worker import claim, run_one


@pytest.fixture
def client():
    with TestClient(app) as c:
        email = f'{time.time_ns()}@example.com'
        assert (
            c.post(
                '/api/auth/register', json={'email': email, 'password': 'a-long-password'}
            ).status_code
            == 200
        )
        yield c


def wid(c):
    return c.get('/api/workspaces').json()[0]['id']


def upload(
    c,
    w,
    name='brief.md',
    text='The approved project budget is INR 75000. Maya owns the launch on 15 December 2026.',
):
    response = c.post(
        f'/api/workspaces/{w}/documents', files=[('files', (name, text.encode(), 'text/plain'))]
    )
    assert response.status_code == 200
    return response.json()['documents'][0]


def drain():
    for _ in range(30):
        if not run_one():
            break


def test_upload_retrieve_deduplicate_delete(client):
    w = wid(client)
    d = upload(client, w)
    assert d['status'] == 'queued'
    duplicate = upload(client, w)
    assert duplicate['id'] == d['id'] and duplicate['duplicate']
    drain()
    docs = client.get(f'/api/workspaces/{w}/documents').json()
    assert docs[0]['status'] == 'ready' and docs[0]['chunk_count'] > 0
    answer = client.post(
        f'/api/workspaces/{w}/chat', json={'question': 'What is the project budget?'}
    ).json()
    assert answer['mode'] == 'local_excerpts'
    assert '75000' in answer['answer']
    assert answer['sources'][0]['document_id'] == d['id']
    preview = client.get(f'/api/workspaces/{w}/documents/{d["id"]}/chunks').json()
    assert preview['chunks'][0]['text'] == answer['sources'][0]['text']
    assert client.delete(f'/api/workspaces/{w}/documents/{d["id"]}').status_code == 200
    assert not (config.DATA / 'files' / d['id']).exists()
    answer = client.post(f'/api/workspaces/{w}/chat', json={'question': 'budget'}).json()
    assert answer['mode'] == 'insufficient_evidence' and not answer['sources']


def test_accounts_and_workspace_isolation(client):
    w = wid(client)
    d = upload(client, w)
    drain()
    with TestClient(app) as other:
        assert other.get(f'/api/workspaces/{w}/documents').status_code == 401
        other.post(
            '/api/auth/register',
            json={'email': f'{time.time_ns()}@other.com', 'password': 'another-long-password'},
        )
        for suffix in [
            '/documents',
            f'/documents/{d["id"]}/chunks',
            f'/documents/{d["id"]}/original',
        ]:
            assert other.get(f'/api/workspaces/{w}{suffix}').status_code == 404
        assert other.delete(f'/api/workspaces/{w}/documents/{d["id"]}').status_code == 404
        other_w = wid(other)
        assert (
            other.post(
                f'/api/workspaces/{other_w}/chat',
                json={'question': 'budget', 'document_ids': [d['id']]},
            ).status_code
            == 404
        )
        answer = other.post(f'/api/workspaces/{other_w}/chat', json={'question': 'budget'}).json()
        assert not answer['sources']
    assert client.post('/api/auth/logout').status_code == 200
    assert client.get('/api/me').status_code == 401


def test_processing_failure_retry_and_lease_recovery(client):
    w = wid(client)
    d = upload(client, w, 'broken.pdf', 'not a PDF')
    drain()
    docs = client.get(f'/api/workspaces/{w}/documents').json()
    assert docs[0]['status'] == 'failed'
    assert (
        client.post(f'/api/workspaces/{w}/documents/{d["id"]}/retry').json()['status'] == 'queued'
    )
    claimed = claim(Document)
    assert claimed[0] == d['id']
    assert claim(Document) is None
    with Session() as db:
        doc = db.get(Document, d['id'])
        doc.lease = time.time() - 1
        db.commit()
    reclaimed = claim(Document)
    assert reclaimed == (d['id'], claimed[1] + 1)
    drain()


def test_scoped_search_and_unanswerable(client):
    w = wid(client)
    first = upload(client, w, 'finance.md', 'The finance budget is 75000 rupees.')
    second = upload(client, w, 'travel.md', 'Travel reimbursement requires manager approval.')
    drain()
    answer = client.post(
        f'/api/workspaces/{w}/chat', json={'question': 'budget', 'document_ids': [second['id']]}
    ).json()
    assert answer['mode'] == 'insufficient_evidence'
    answer = client.post(
        f'/api/workspaces/{w}/chat', json={'question': 'budget', 'document_ids': [first['id']]}
    ).json()
    assert {s['document_id'] for s in answer['sources']} == {first['id']}
    assert client.post(f'/api/workspaces/{w}/chat', json={'question': '   '}).status_code == 422


def test_validation_and_password_login(client):
    w = wid(client)
    bad = upload(client, w, 'virus.exe', 'text')
    assert 'error' in bad and 'id' not in bad
    assert 'error' in upload(client, w, 'empty.txt', '')
    email = client.get('/api/me').json()['email']
    client.post('/api/auth/logout')
    assert (
        client.post(
            '/api/auth/login', json={'email': email, 'password': 'wrong-long-password'}
        ).status_code
        == 401
    )
    assert (
        client.post(
            '/api/auth/login', json={'email': email, 'password': 'a-long-password'}
        ).status_code
        == 200
    )
    assert client.get('/api/me').status_code == 200


def test_chunker_preserves_end_and_bounded_size():
    original = 'word ' * 2000 + 'END_MARKER'
    chunks = list(split_text(original))
    assert chunks[-1].endswith('END_MARKER')
    assert all(len(c) <= 1800 for c in chunks)


def test_ai_invalid_citations_rejected(monkeypatch):
    monkeypatch.setattr(config, 'API_KEY', 'test-key')
    fake = SimpleNamespace(
        responses=SimpleNamespace(
            create=lambda **kwargs: SimpleNamespace(output_text='Budget is 99 [7]')
        )
    )
    monkeypatch.setattr(ai, 'client', lambda: fake)
    result = ai.generate('budget', [{'name': 'test', 'page': 1, 'text': 'Budget is 20'}])
    assert result['mode'] == 'citation_check_failed'
    fake.responses.create = lambda **kwargs: SimpleNamespace(output_text='Budget is 20 [1]')
    assert (
        ai.generate('budget', [{'name': 'test', 'page': 1, 'text': 'Budget is 20'}])['mode'] == 'ai'
    )


def test_ai_embedding_order(monkeypatch):
    monkeypatch.setattr(config, 'API_KEY', 'test-key')
    fake = SimpleNamespace(
        embeddings=SimpleNamespace(
            create=lambda **kwargs: SimpleNamespace(
                data=[
                    SimpleNamespace(index=1, embedding=[2]),
                    SimpleNamespace(index=0, embedding=[1]),
                ]
            )
        )
    )
    monkeypatch.setattr(ai, 'client', lambda: fake)
    assert ai.embed(['one', 'two']) == [[1], [2]]


def test_pdf_page_sources_and_blank_page_rejection(client):
    import io

    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject(
        {
            NameObject('/Type'): NameObject('/Font'),
            NameObject('/Subtype'): NameObject('/Type1'),
            NameObject('/BaseFont'): NameObject('/Helvetica'),
        }
    )
    page[NameObject('/Resources')] = DictionaryObject(
        {NameObject('/Font'): DictionaryObject({NameObject('/F1'): font})}
    )
    stream = DecodedStreamObject()
    stream.set_data(b'BT /F1 12 Tf 50 700 Td (The approved budget is 12345 rupees.) Tj ET')
    page[NameObject('/Contents')] = writer._add_object(stream)
    out = io.BytesIO()
    writer.write(out)
    w = wid(client)
    response = client.post(
        f'/api/workspaces/{w}/documents',
        files={'files': ('text.pdf', out.getvalue(), 'application/pdf')},
    )
    assert response.status_code == 200
    drain()
    result = client.post(f'/api/workspaces/{w}/chat', json={'question': 'budget'}).json()
    assert result['sources'][0]['page'] == 1
    assert result['sources'][0]['location_type'] == 'page'
    assert '12345' in result['answer']
    writer.add_blank_page(width=612, height=792)
    out = io.BytesIO()
    writer.write(out)
    client.post(
        f'/api/workspaces/{w}/documents',
        files={'files': ('partial-scan.pdf', out.getvalue(), 'application/pdf')},
    )
    drain()
    doc = next(
        d
        for d in client.get(f'/api/workspaces/{w}/documents').json()
        if d['name'] == 'partial-scan.pdf'
    )
    assert doc['status'] == 'failed' and 'OCR' in doc['error']
    assert doc['chunk_count'] == 0


def test_docx_tables_and_sections(client):
    import io

    from docx import Document as WordDocument

    doc = WordDocument()
    doc.add_paragraph('Launch owner is Maya.')
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = 'Budget'
    table.cell(0, 1).text = '45000'
    out = io.BytesIO()
    doc.save(out)
    w = wid(client)
    client.post(
        f'/api/workspaces/{w}/documents',
        files={
            'files': (
                'brief.docx',
                out.getvalue(),
                'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
            )
        },
    )
    drain()
    result = client.post(f'/api/workspaces/{w}/chat', json={'question': 'budget'}).json()
    assert result['sources'][0]['location_type'] == 'section'
    assert '45000' in result['answer']


def test_upload_limits(client, monkeypatch):
    from app import main

    w = wid(client)
    upload(client, w)
    drain()
    upload(client, w, 'new.md', 'New evidence changes the budget to 80000.')
    monkeypatch.setattr(main, 'MAX_FILE_BYTES', 10)
    result = upload(client, w, 'too-large.md', 'x' * 11)
    assert 'error' in result and 'id' not in result
    files = list((config.DATA / 'files').iterdir())
    with Session() as db:
        ids = {d.id for d in db.query(Document)}
    assert all(p.name in ids for p in files)


def test_hybrid_semantic_retrieval(client, monkeypatch):
    from app import retrieval

    w = wid(client)
    d = upload(client, w, 'retention.md', 'Customer churn fell by ten percent.')
    drain()
    vector = [1.0] + [0.0] * 1535
    with Session() as db:
        chunk = db.query(Chunk).filter_by(document_id=d['id']).first()
        chunk.embedding = vector
        chunk.embedding_model = config.EMBED_MODEL
        db.commit()
    monkeypatch.setattr(config, 'API_KEY', 'mock-provider')
    monkeypatch.setattr(retrieval, 'embed', lambda texts: [vector for _ in texts])
    # No literal overlap; the mocked embedding supplies the semantic match.
    sources = retrieval.retrieve(w, 'subscriber attrition')
    assert sources[0]['document_id'] == d['id']


def test_stale_job_cannot_publish_after_reindex(client):
    from app.ingest import ingest

    w = wid(client)
    upload(client, w)
    job_id, generation = claim(Document)
    with Session() as db:
        doc = db.get(Document, job_id)
        doc.generation += 1
        doc.status = 'queued'
        db.commit()
    ingest(job_id, generation, lambda: True)
    with Session() as db:
        assert db.get(Document, job_id).status == 'queued'
        assert db.query(Chunk).filter_by(document_id=job_id).count() == 0
    drain()


def test_local_chat_is_question_specific_and_topic_changes_do_not_reuse_history(client):
    w = wid(client)
    text = (
        '# Pilot brief\nThe monthly infrastructure budget is INR 75000. '
        'The Atlas pilot launches on 15 December 2026. '
        'Maya owns the pilot launch. '
        'Scanned PDFs require OCR before uploading.'
    )
    upload(client, w, 'pilot.md', text)
    drain()
    history = []

    def ask(question):
        result = client.post(
            f'/api/workspaces/{w}/chat', json={'question': question, 'history': history}
        ).json()
        history.extend(
            [
                {'role': 'user', 'content': question},
                {'role': 'assistant', 'content': result['answer']},
            ]
        )
        return result

    budget = ask('infra budget')
    assert '75000' in budget['answer']
    # Local excerpt mode exposes complete retrieved passages, not a focused AI answer.
    launch = ask('pilot launch date')
    assert '15 December 2026' in launch['answer']
    unrelated = ask('leave policy')
    assert unrelated['mode'] == 'insufficient_evidence' and unrelated['sources'] == []
    repeated = ask('leave policy')
    assert repeated == unrelated
    # Source preview still contains the complete original evidence.
    assert budget['sources'][0]['text'] == text


def test_explicit_followup_retains_context_without_contaminating_new_topics():
    from app.query_context import retrieval_query

    history = [SimpleNamespace(role='user', content='Tell me about the Atlas pilot')]
    assert retrieval_query('What is its budget?', history).endswith('Tell me about the Atlas pilot')
    assert retrieval_query('leave policy', history) == 'leave policy'
    assert retrieval_query('infra budget', history) == 'infra budget'


def test_chunked_upload_offsets_completion_duplicates_and_access(client):
    w = wid(client)
    data = b'A large document has a project budget of 12345.\n' * 20
    start = client.post(
        f'/api/workspaces/{w}/uploads', json={'name': 'large.md', 'size': len(data)}
    )
    assert start.status_code == 200
    job = start.json()['id']
    base = f'/api/workspaces/{w}/uploads/{job}'
    assert client.post(base + '/complete').status_code == 409
    first = client.put(
        base + '?offset=0', content=data[:100], headers={'Content-Type': 'application/octet-stream'}
    )
    assert first.json()['offset'] == 100
    assert client.put(base + '?offset=0', content=data[:100]).status_code == 409
    assert client.get(base).json()['offset'] == 100
    assert client.put(base + '?offset=100', content=data[100:]).status_code == 200
    assert client.put(base + f'?offset={len(data)}', content=b'extra').status_code == 413
    with TestClient(app) as other:
        other.post(
            '/api/auth/register',
            json={'email': f'{time.time_ns()}@upload.test', 'password': 'another-long-password'},
        )
        ow = wid(other)
        assert other.get(base).status_code == 404
        assert other.get(f'/api/workspaces/{ow}/uploads/{job}').status_code == 404
        assert (
            other.put(f'/api/workspaces/{ow}/uploads/{job}?offset=0', content=b'x').status_code
            == 404
        )
    result = client.post(base + '/complete')
    assert result.status_code == 200 and not result.json()['duplicate']
    assert client.get(base).status_code == 404
    drain()
    d = result.json()
    assert (config.DATA / 'files' / d['id']).read_bytes() == data
    duplicate = client.post(
        f'/api/workspaces/{w}/uploads', json={'name': 'same.md', 'size': len(data)}
    ).json()['id']
    path = f'/api/workspaces/{w}/uploads/{duplicate}'
    client.put(path + '?offset=0', content=data)
    completed = client.post(path + '/complete').json()
    assert completed['duplicate'] and completed['id'] == d['id']


def test_staged_upload_validation_and_cancellation(client, monkeypatch):
    from app import main

    w = wid(client)
    monkeypatch.setattr(main, 'MAX_FILE_BYTES', 100)
    path = f'/api/workspaces/{w}/uploads'
    assert client.post(path, json={'name': 'big.pdf', 'size': 101}).status_code == 413
    assert client.post(path, json={'name': 'bad.exe', 'size': 10}).status_code == 422
    job = client.post(path, json={'name': 'valid.md', 'size': 50}).json()['id']
    assert client.delete(path + '/' + job).status_code == 200
    assert client.get(path + '/' + job).status_code == 404
    assert not (config.DATA / 'uploads' / f'{job}.part').exists()


def test_country_count_answers_numeric_evidence_instead_of_titles(client):
    w = wid(client)
    upload(
        client,
        w,
        'Asia.md',
        '## World map\nAsia North America\nNorth America\nA continent in many voices',
    )
    data = (
        'NORTH AMERICA / ABOUT THIS VOLUME\nA continent in many voices\n'
        'This book uses a broad geographic definition of North America, including Central America and the Caribbean. '
        'It gives each of the region’s 23 sovereign countries its own chapter. '
        'Greenland and Puerto Rico are non-sovereign places in a separate reference section.'
    )
    target = upload(client, w, 'North_America.md', data)
    drain()
    result = client.post(
        f'/api/workspaces/{w}/chat',
        json={'question': 'how many countries are there in North America'},
    ).json()
    assert result['mode'] == 'local_excerpts'
    assert '23 sovereign countries' in result['answer']
    assert any(source['document_id'] == target['id'] for source in result['sources'])


def test_saved_conversations_are_scoped_and_editable(client):
    w = wid(client)
    route = f'/api/workspaces/{w}/conversations'
    first = client.post(
        route,
        json={
            'title': 'Budget research',
            'messages': [{'role': 'user', 'content': 'What is the budget?'}],
        },
    )
    assert first.status_code == 200
    cid = first.json()['id']
    assert client.get(route).json()[0]['id'] == cid
    assert client.get(f'{route}/{cid}').json()['messages'][0]['content'] == 'What is the budget?'
    updated = client.put(
        f'{route}/{cid}',
        json={
            'title': 'Budget research',
            'messages': [
                {'role': 'user', 'content': 'What is the budget?'},
                {
                    'role': 'assistant',
                    'content': 'INR 75000 [1]',
                    'result': {'answer': 'INR 75000 [1]', 'sources': [], 'mode': 'ai'},
                },
            ],
        },
    )
    assert updated.status_code == 200
    assert len(updated.json()['messages']) == 2
    renamed = client.patch(f'{route}/{cid}', json={'title': 'Budget notes'})
    assert renamed.status_code == 200 and renamed.json()['title'] == 'Budget notes'
    with TestClient(app) as other:
        other.post(
            '/api/auth/register',
            json={'email': f'{time.time_ns()}@second.com', 'password': 'another-long-password'},
        )
        ow = wid(other)
        assert other.get(f'/api/workspaces/{ow}/conversations/{cid}').status_code == 404
        assert other.delete(f'/api/workspaces/{ow}/conversations/{cid}').status_code == 404
    assert client.delete(f'{route}/{cid}').status_code == 200
    assert client.get(f'{route}/{cid}').status_code == 404


def test_insights_endpoints_are_removed(client):
    w = wid(client)
    assert client.get(f'/api/workspaces/{w}/analyses').status_code == 404
    assert client.post(f'/api/workspaces/{w}/analyses', json={'kind': 'summary'}).status_code == 404


def test_workspace_deletion_requires_owner_and_confirmation(client):
    w = wid(client)
    document = upload(client, w)
    drain()
    created = client.post(
        f'/api/workspaces/{w}/conversations', json={'title': 'Delete test', 'messages': []}
    ).json()
    assert (
        client.request('DELETE', f'/api/workspaces/{w}', json={'name': 'wrong'}).status_code == 422
    )
    with TestClient(app) as other:
        other.post(
            '/api/auth/register',
            json={'email': f'{time.time_ns()}@another.com', 'password': 'another-long-password'},
        )
        assert (
            other.request(
                'DELETE', f'/api/workspaces/{w}', json={'name': 'My knowledge base'}
            ).status_code
            == 404
        )
    result = client.request('DELETE', f'/api/workspaces/{w}', json={'name': 'My knowledge base'})
    assert result.status_code == 200
    assert not (config.DATA / 'files' / document['id']).exists()
    assert client.get(f'/api/workspaces/{w}/conversations/{created["id"]}').status_code == 404
    assert all(item['id'] != w for item in client.get('/api/workspaces').json())


def test_gmail_routes_are_authenticated_and_config_guarded(client, monkeypatch):
    from app import gmail

    w = wid(client)
    route = f'/api/workspaces/{w}/gmail'
    assert client.get(f'/api/gmail/status?wid={w}').json()['connected'] is False
    assert client.post(route + '/search', json={'query': 'in:inbox'}).status_code == 409
    assert client.post(route + '/summarize', json={'message_ids': []}).status_code == 422
    with TestClient(app) as stranger:
        assert stranger.get(f'/api/gmail/status?wid={w}').status_code == 401
    monkeypatch.setattr(gmail, 'CLIENT', '')
    assert client.get(f'/api/gmail/connect?wid={w}').status_code == 503


def test_gmail_summary_reports_provider_configuration_without_gmail_access(client, monkeypatch):
    from app import gmail

    w = wid(client)
    monkeypatch.setattr(gmail.config, 'AI_PROVIDER', 'none')
    monkeypatch.setattr(gmail.config, 'API_KEY', '')
    result = client.post(f'/api/workspaces/{w}/gmail/summarize', json={'message_ids': ['abc123']})
    assert result.status_code == 409
    assert 'AI provider' in result.json()['detail']


def test_gmail_email_text_decoder():
    import base64

    from app.gmail import text_body

    original = 'A deadline is next Tuesday.'
    encoded = base64.urlsafe_b64encode(original.encode()).decode().rstrip('=')
    assert (
        text_body(
            {
                'mimeType': 'multipart/alternative',
                'parts': [{'mimeType': 'text/plain', 'body': {'data': encoded}}],
            }
        )
        == original
    )


def test_connected_chat_requires_sign_in():
    with TestClient(app) as unauthenticated:
        response = unauthenticated.post(
            '/api/workspaces/unowned/connected-chat', json={'question': 'Summarize my sources'}
        )
        assert response.status_code == 401


def test_connected_chat_rejects_invalid_gmail_scope(client):
    import pytest

    from app.main import ConnectedQuestion

    with pytest.raises(ValueError):
        ConnectedQuestion(question='What changed?', include_gmail=True, gmail_query='x' * 401)


def test_connected_source_evidence_contract():
    from app.sources import Evidence

    source = Evidence(
        'gmail', 'message-id', 'Update', 'The launch was moved.', 'https://mail.google.com/'
    )
    prompt = source.as_prompt(4)
    assert prompt['source_id'] == 4
    assert prompt['provider'] == 'gmail'
    assert prompt['title'] == 'Update'


def test_notion_source_search_reads_shared_page_content(monkeypatch):
    from app import external_sources

    page_id = '29601486-3e4f-8030-b34b-d03b4a4f6944'
    calls = []

    def fake_http_json(url, token, provider, params=None, method='GET'):
        calls.append((url, params, method))
        if url.endswith('/v1/search') and params and params.get('query'):
            return {'results': [], 'has_more': False}
        if url.endswith('/v1/search'):
            return {
                'results': [
                    {
                        'object': 'page',
                        'id': page_id,
                        'url': 'https://notion.so/hackathon',
                        'properties': {'title': {'title': [{'plain_text': 'Hackathon'}]}},
                    }
                ]
            }
        if url.endswith(f'/v1/blocks/{page_id}/children'):
            return {
                'results': [
                    {
                        'id': 'block-1',
                        'type': 'paragraph',
                        'has_children': False,
                        'paragraph': {
                            'rich_text': [
                                {'plain_text': 'Organizers assign roles to judges and mentors.'}
                            ]
                        },
                    }
                ],
                'has_more': False,
            }
        raise AssertionError(f'Unexpected Notion API request: {url}')

    monkeypatch.setattr(external_sources, 'http_json', fake_http_json)

    results = external_sources.source_search('notion', 'test-token', 'organizer roles')

    assert len(results) == 1
    assert results[0]['id'] == page_id
    assert results[0]['title'] == 'Hackathon'
    assert 'Organizers assign roles' in results[0]['excerpt']
    assert any('/blocks/' in url for url, _, _ in calls)


def test_slack_bot_token_error_explains_user_search_token_requirement(monkeypatch):
    from fastapi import HTTPException

    from app import external_sources

    monkeypatch.setattr(
        external_sources,
        'http_json',
        lambda *args, **kwargs: {'ok': False, 'error': 'not_allowed_token_type'},
    )
    with pytest.raises(HTTPException) as exc:
        external_sources.source_search('slack', 'test-token', 'notion')
    assert exc.value.status_code == 422
    assert 'user-scoped token' in exc.value.detail
    assert 'search:read' in exc.value.detail


def test_slack_search_keeps_long_message_text_for_relevant_evidence(monkeypatch):
    from app import external_sources

    full_text = (
        'Deforestation note. '
        + ('Background text. ' * 180)
        + 'Causes: agricultural expansion, urbanisation, industrialisation, timber logging, mining, and forest fires.'
    )
    monkeypatch.setattr(
        external_sources,
        'http_json',
        lambda *args, **kwargs: {
            'ok': True,
            'messages': {
                'matches': [
                    {
                        'ts': '1',
                        'channel': {'name': 'environment'},
                        'text': full_text,
                        'permalink': 'https://slack.test/message',
                    }
                ]
            },
        },
    )
    result = external_sources.source_search('slack', 'token', 'causes of deforestation')
    assert result[0]['excerpt'] == full_text
    assert 'forest fires' in result[0]['excerpt']


def test_slack_search_uses_message_text_when_channel_label_is_an_id(monkeypatch):
    from app import external_sources

    monkeypatch.setattr(
        external_sources,
        'http_json',
        lambda *args, **kwargs: {
            'ok': True,
            'messages': {
                'matches': [
                    {
                        'ts': '1',
                        'channel': {'name': 'U0CDTS8H4Q2'},
                        'text': 'Causes include agricultural expansion. More context follows.',
                        'permalink': 'https://slack.test/message',
                    }
                ]
            },
        },
    )

    result = external_sources.source_search('slack', 'token', 'causes deforestation')

    assert result[0]['title'] == 'Causes include agricultural expansion'


def test_connected_payload_keeps_relevant_late_slack_details_and_ranks_them_first():
    from app.connected_synthesis import bounded_payload
    from app.sources import Evidence

    slack_text = (
        'Deforestation overview. '
        + ('General background text. ' * 180)
        + (
            'Causes include agricultural expansion, urbanisation, industrialisation, '
            'timber logging, mining activities, and forest fires.'
        )
    )
    evidence = [
        Evidence('documents', 'doc', 'General guide', 'Deforestation is a problem.'),
        Evidence('slack', 'msg', '#environment', slack_text),
    ]
    payload = bounded_payload(evidence, question='What are the causes of deforestation?')
    assert payload[0]['provider'] == 'slack'
    assert all(
        term in payload[0]['text'].lower()
        for term in (
            'agricultural expansion',
            'urbanisation',
            'industrialisation',
            'timber logging',
            'mining activities',
            'forest fires',
        )
    )


def test_connected_synthesis_fits_local_model_and_preserves_citations():
    from app.connected_synthesis import synthesize_connected
    from app.sources import Evidence

    evidence = [Evidence('documents', 'd', 'Document', 'important fact ' * 1000)]
    evidence += [Evidence('gmail', str(i), 'Mail', 'other fact ' * 1000) for i in range(8)]
    calls = []

    def fake_ollama(messages):
        import json

        payload = json.loads(messages[-1]['content'])
        calls.append(payload)
        if len(calls) == 1:
            raise ValueError('Model response was truncated; try a narrower question.')
        return 'An important fact is in the document [1]. Other fact is present in email [2].'

    answer, payload = synthesize_connected(
        'What changed?', evidence, 'ollama', '', 'ignored', None, fake_ollama
    )
    assert len(calls) == 2
    assert len(calls[0]['sources']) <= 50
    assert sum(len(x['text']) for x in calls[0]['sources']) <= 16000
    assert [x['source_id'] for x in payload] == list(range(1, len(payload) + 1))
    assert {x['provider'] for x in payload} == {'documents', 'gmail'}
    assert answer == 'An important fact is in the document [1]. Other fact is present in email [2].'


def test_connected_synthesis_includes_broad_repository_inventory():
    from app.connected_synthesis import bounded_payload
    from app.sources import Evidence

    evidence = [
        Evidence(
            'github',
            str(i),
            f'owner/project-{i}',
            f'Project {i} description. Language: Python. Topics: search, research.',
        )
        for i in range(37)
    ]
    payload = bounded_payload(evidence)
    assert len(payload) == 37
    assert [item['origin_id'] for item in payload] == [str(i) for i in range(37)]


def test_connected_synthesis_hides_reasoning_and_validates_before_emitting(monkeypatch):
    from contextlib import contextmanager

    from app import connected_synthesis
    from app.sources import Evidence

    lines = [
        {'message': {'content': 'Let me analyze the source first. '}},
        {'message': {'content': 'The change adds a retry [1].'}},
        {'done': True, 'done_reason': 'stop'},
    ]

    class Response:
        def raise_for_status(self):
            pass

        def iter_lines(self):
            import json

            return iter(json.dumps(line) for line in lines)

    @contextmanager
    def fake_stream(*args, **kwargs):
        yield Response()

    monkeypatch.setattr(connected_synthesis.httpx, 'stream', fake_stream)
    events = list(
        connected_synthesis.stream_connected(
            'Summarize the change',
            [Evidence('github', '1', 'Retry PR', 'Adds retry support')],
            'ollama',
            '',
            'model',
            None,
            'http://ollama',
            'model',
        )
    )
    assert [event['type'] for event in events] == ['complete']
    assert events[0]['mode'] == 'generation_failed'
    assert 'Let me analyze' not in events[0]['answer']
    assert events[0]['sources'] == []


def test_connected_synthesis_emits_only_citation_checked_final_text(monkeypatch):
    from contextlib import contextmanager

    from app import connected_synthesis
    from app.sources import Evidence

    lines = [
        {'message': {'content': 'The change adds a retry [1].'}},
        {'done': True, 'done_reason': 'stop'},
    ]

    class Response:
        def raise_for_status(self):
            pass

        def iter_lines(self):
            import json

            return iter(json.dumps(line) for line in lines)

    @contextmanager
    def fake_stream(*args, **kwargs):
        yield Response()

    monkeypatch.setattr(connected_synthesis.httpx, 'stream', fake_stream)
    events = list(
        connected_synthesis.stream_connected(
            'Summarize the change',
            [Evidence('github', '1', 'Retry PR', 'Adds retry support')],
            'ollama',
            '',
            'model',
            None,
            'http://ollama',
            'model',
        )
    )
    assert [event['type'] for event in events] == ['delta', 'complete']
    assert events[0]['text'] == 'The change adds a retry [1].'
    assert events[1]['mode'] == 'connected'


def test_local_connected_synthesis_timeout_does_not_masquerade_as_answer(monkeypatch):
    from contextlib import contextmanager

    from app import connected_synthesis
    from app.sources import Evidence

    class Response:
        def raise_for_status(self):
            pass

        def iter_lines(self):
            return iter(['{"message":{"content":"Still reasoning…"}}'])

    @contextmanager
    def fake_stream(*args, **kwargs):
        yield Response()

    ticks = iter([0, 0, 241])
    monkeypatch.setattr(connected_synthesis.httpx, 'stream', fake_stream)
    monkeypatch.setattr(connected_synthesis.time, 'monotonic', lambda: next(ticks))
    events = list(
        connected_synthesis.stream_connected(
            'What is deforestation?',
            [
                Evidence(
                    'documents', '1', 'Forest guide', 'Deforestation removes native forest cover.'
                )
            ],
            'ollama',
            '',
            'model',
            None,
            'http://ollama',
            'model',
        )
    )
    assert events[-1]['type'] == 'complete' and events[-1]['mode'] == 'generation_failed'
    assert 'timed out' in events[-1]['answer']
    assert events[-1]['sources'] == []


def test_connected_chat_stream_returns_explicit_no_evidence_result(client):
    w = wid(client)
    response = client.post(
        f'/api/workspaces/{w}/connected-chat/stream',
        json={'question': 'Summarize the selected sources', 'providers': []},
    )
    assert response.status_code == 200
    events = [__import__('json').loads(line) for line in response.text.splitlines() if line]
    assert events[0]['type'] == 'status'
    assert 'Documents' in events[0]['text']
    assert events[-1]['type'] == 'done'
    assert events[-1]['result']['mode'] == 'insufficient_evidence'
    assert events[-1]['result']['retrieval']['passages_considered'] == 0
    assert events[-1]['result']['retrieval']['sources'] == [
        {'provider': 'documents', 'status': 'searched', 'found': 0, 'considered': 0, 'used': 0}
    ]


def test_connected_chat_stream_keeps_local_excerpt_mode_without_model(client, monkeypatch):
    w = wid(client)
    upload(client, w)
    drain()
    monkeypatch.setattr(config, 'AI_PROVIDER', 'none')
    monkeypatch.setattr(config, 'API_KEY', '')
    response = client.post(
        f'/api/workspaces/{w}/connected-chat/stream',
        json={'question': 'What is the budget?', 'providers': []},
    )
    assert response.status_code == 200
    events = [__import__('json').loads(line) for line in response.text.splitlines() if line]
    assert events[-1]['type'] == 'done'
    assert events[-1]['result']['mode'] == 'local_excerpts'
    assert '75000' in events[-1]['result']['answer']


def test_connected_chat_can_exclude_documents_and_reports_source_specific_no_match(
    client, monkeypatch
):
    from app import external_sources, main
    from app.db import ExternalConnection, Session

    w = wid(client)
    with Session() as db:
        db.add(
            ExternalConnection(workspace_id=w, provider='slack', account='test', secret='encrypted')
        )
        db.commit()
    monkeypatch.setattr(
        main,
        'retrieve',
        lambda *args, **kwargs: (_ for _ in ()).throw(
            AssertionError('documents should be excluded')
        ),
    )
    monkeypatch.setattr(external_sources, 'read_connection', lambda *args: (None, 'test-token'))
    monkeypatch.setattr(external_sources, 'source_search', lambda *args: [])
    response = client.post(
        f'/api/workspaces/{w}/connected-chat/stream',
        json={
            'question': 'What did we decide in Slack?',
            'include_documents': False,
            'providers': ['slack'],
        },
    )
    assert response.status_code == 200
    events = [__import__('json').loads(line) for line in response.text.splitlines() if line]
    result = events[-1]['result']
    assert result['mode'] == 'insufficient_evidence'
    assert result['scope'] == ['slack']
    assert result['retrieval']['sources'] == [
        {'provider': 'slack', 'status': 'searched', 'found': 0, 'considered': 0, 'used': 0}
    ]
    assert 'No matching results were found in Slack' in result['answer']


def test_connected_chat_preserves_full_slack_message_through_retrieval(client, monkeypatch):
    from app import connected_synthesis, external_sources, main

    w = wid(client)
    full_text = (
        'Deforestation introduction. '
        + ('Background. ' * 150)
        + (
            'Causes include agricultural expansion, urbanisation, industrialisation, '
            'timber logging, mining activities, and forest fires.'
        )
    )
    monkeypatch.setattr(main, 'retrieve', lambda *args, **kwargs: [])
    monkeypatch.setattr(external_sources, 'read_connection', lambda *args: (None, 'token'))
    monkeypatch.setattr(
        external_sources,
        'source_search',
        lambda *args: [
            {
                'id': 'slack-1',
                'title': '#environment',
                'excerpt': full_text,
                'url': 'https://slack.test/1',
            }
        ],
    )
    captured = {}

    def fake_synthesize(question, evidence, *args):
        item = evidence[0]
        captured['excerpt'] = item.excerpt
        return 'The note lists several causes [1].', [
            {**item.as_prompt(1), 'origin_id': item.source_id}
        ]

    monkeypatch.setattr(connected_synthesis, 'synthesize_connected', fake_synthesize)
    response = client.post(
        f'/api/workspaces/{w}/connected-chat',
        json={
            'question': 'What causes of deforestation are listed?',
            'include_documents': False,
            'providers': ['slack'],
        },
    )
    assert response.status_code == 200, response.text
    assert captured['excerpt'] == full_text
    assert 'forest fires' in captured['excerpt']


def test_natural_language_connected_queries_are_simplified():
    from app.main import drive_query_for_question, gmail_query_for_question

    assert (
        gmail_query_for_question(
            'newer_than:30d', 'What are the latest project updates in my email?'
        )
        == 'newer_than:30d project updates'
    )
    assert (
        gmail_query_for_question('newer_than:30d', 'Are there any credit card offers?')
        == 'newer_than:30d credit card offers'
    )
    assert (
        drive_query_for_question(
            'Find the AI Engineer resume files in my Google Drive and list their exact file names.'
        )
        == 'ai engineer resume'
    )


def test_document_retrieval_requires_multiple_concept_matches_for_specific_questions():
    from app.retrieval import has_query_support, minimum_query_overlap

    question = 'What is the recipe for matcha ice cream in my documents?'
    assert minimum_query_overlap(question) == 2
    assert not has_query_support(question, 'Iceland, land of fire and ice.')
    assert not has_query_support(question, 'A sample recipe for a project idea.')
    assert has_query_support(
        'What causes of deforestation are listed?',
        'Causes of deforestation include agricultural expansion and logging.',
    )


def test_credit_card_offer_retrieval_drops_unrelated_results_and_reads_full_email(
    client, monkeypatch
):
    import base64

    from app import gmail, main
    from app.db import Session, User

    w = wid(client)
    actor_id = client.get('/api/me').json()['id']
    question = 'Are there any credit card offers?'
    detail = (
        'Your monthly statement is ready. '
        + ('Account information. ' * 100)
        + ('Exclusive credit card cashback offer: earn 5% cashback on dining through December.')
    )
    encoded = base64.urlsafe_b64encode(detail.encode()).decode().rstrip('=')
    messages = {
        'statement': {
            'id': 'statement',
            'payload': {
                'headers': [{'name': 'Subject', 'value': 'SBI Card ELITE Monthly Statement'}],
                'mimeType': 'text/plain',
                'body': {'data': encoded},
            },
        },
        'unrelated': {
            'id': 'unrelated',
            'payload': {
                'headers': [{'name': 'Subject', 'value': 'North America'}],
                'mimeType': 'text/plain',
                'body': {
                    'data': base64.urlsafe_b64encode(
                        b'Country guide and travel information'
                    ).decode()
                },
            },
        },
    }
    monkeypatch.setattr(gmail, 'connection', lambda *args: object())
    monkeypatch.setattr(gmail, 'access_token', lambda *args: 'test-token')
    monkeypatch.setattr(
        gmail,
        'gmail_get',
        lambda token, url, params=None: (
            {'messages': [{'id': 'statement'}, {'id': 'unrelated'}]}
            if url.endswith('/messages')
            else messages[url.rsplit('/', 1)[-1]]
        ),
    )
    body = main.ConnectedQuestion(question=question, include_documents=False, include_gmail=True)
    with Session() as db:
        actor = db.get(User, actor_id)
        evidence, _ = main.collect_connected_evidence(w, body, actor, db)

    assert len(evidence) == 1
    assert evidence[0].provider == 'gmail'
    assert evidence[0].source_id == 'statement'
    assert '5% cashback' in evidence[0].excerpt
    assert len(evidence[0].excerpt) > 1300


def test_github_pull_request_search_honors_requested_count(monkeypatch):
    from app import external_sources

    captured = {}

    def fake_http(url, token, provider, params=None, method='GET'):
        if url.endswith('/user'):
            return {'login': 'test-user'}
        captured.update(params or {})
        return {'items': []}

    monkeypatch.setattr(external_sources, 'http_json', fake_http)
    external_sources.source_search(
        'github', 'token', 'What are three of my recently updated open GitHub pull requests?'
    )
    assert captured['per_page'] == 3


def test_github_own_pull_requests_are_scoped_to_authenticated_login(monkeypatch):
    from app import external_sources

    requests = []

    def fake_http(url, token, provider, params=None, method='GET'):
        requests.append((url, params or {}))
        if url.endswith('/user'):
            return {'login': 'srbmaury'}
        return {
            'items': [
                {
                    'id': 1,
                    'title': 'My PR',
                    'body': 'Update',
                    'html_url': 'https://github.com/org/repo/pull/1',
                }
            ]
        }

    monkeypatch.setattr(external_sources, 'http_json', fake_http)
    result = external_sources.source_search(
        'github', 'token', 'List three of my recently updated open GitHub pull requests'
    )
    assert requests[0][0].endswith('/user')
    assert requests[1][1]['q'] == 'is:pr state:open author:srbmaury'
    assert requests[1][1]['per_page'] == 3
    assert len(result) == 1


def test_github_own_pull_requests_fail_closed_without_identity(monkeypatch):
    from fastapi import HTTPException

    from app import external_sources

    def fake_http(url, token, provider, params=None, method='GET'):
        return {} if url.endswith('/user') else {'items': []}

    monkeypatch.setattr(external_sources, 'http_json', fake_http)
    with pytest.raises(HTTPException, match='Could not identify'):
        external_sources.source_search('github', 'token', 'List my open pull requests')


def test_github_pull_requests_from_my_github_are_scoped_to_the_connected_author(monkeypatch):
    from app import external_sources

    captured = {}

    def fake_http(url, token, provider, params=None, method='GET'):
        if url.endswith('/user'):
            return {'login': 'srbmaury'}
        captured.update(params or {})
        return {'items': []}

    monkeypatch.setattr(external_sources, 'http_json', fake_http)
    external_sources.source_search('github', 'token', 'List all open pull requests from my GitHub')
    assert captured['q'] == 'is:pr state:open author:srbmaury'


def test_source_health_check_distinguishes_unconnected_from_healthy(client, monkeypatch):
    from app import external_sources
    from app.db import ExternalConnection, Session

    w = wid(client)
    with Session() as db:
        db.add(
            ExternalConnection(
                workspace_id=w, provider='github', account='test-user', secret='encrypted'
            )
        )
        db.commit()
    monkeypatch.setattr(
        external_sources,
        'read_connection',
        lambda db, who, workspaceId, provider: (
            db.get(ExternalConnection, (workspaceId, provider)),
            'test-token',
        ),
    )
    monkeypatch.setattr(
        external_sources, 'verify_source_access', lambda provider, token: ('test-user', 2)
    )
    response = client.post(f'/api/workspaces/{w}/sources/health/check')
    assert response.status_code == 200
    checks = {item['provider']: item for item in response.json()['checks']}
    assert checks['github']['connected'] and checks['github']['healthy']
    assert checks['github']['matches'] == 2
    assert not checks['slack']['connected'] and not checks['slack']['healthy']
    assert 'test-token' not in response.text


def test_github_search_query_preserves_filters_and_simplifies_questions():
    from app.external_sources import (
        github_pull_request_query,
        github_repository_query,
        github_search_query,
        requested_result_limit,
        slack_search_query,
    )

    assert (
        github_search_query('repo:owner/project is:issue bug') == 'repo:owner/project is:issue bug'
    )
    assert github_search_query('repo:owner/project bug') == 'repo:owner/project bug is:issue'
    assert (
        github_search_query('What issues discuss a payment timeout?')
        == 'discuss payment timeout is:issue'
    )
    assert len(github_search_query('what ' * 100 + 'bug')) <= 240
    assert (
        github_pull_request_query('List all open pull requests from my github')
        == 'is:pr state:open'
    )
    assert (
        github_pull_request_query('repo:owner/project is:pull-request state:open')
        == 'repo:owner/project is:pr state:open'
    )
    assert (
        github_repository_query(
            'What different kinds of projects are there in my GitHub?', 'srbmaury'
        )
        == 'user:srbmaury'
    )
    assert (
        requested_result_limit('What are three of my recently updated open GitHub pull requests?')
        == 3
    )
    assert requested_result_limit('List at least 10 of my projects from GitHub') == 10
    assert (
        slack_search_query('What decisions about Hackathon did the team discuss in Slack?')
        == 'Hackathon'
    )


def test_relative_data_dir_is_anchored_to_project_root(tmp_path):
    import os

    from app.config import ROOT, resolve_data_dir

    previous = Path.cwd()
    os.chdir(tmp_path)
    try:
        assert resolve_data_dir('./data', ROOT) == Path(ROOT, 'data').resolve()
        assert resolve_data_dir('/tmp/atlas-data', ROOT) == Path('/tmp/atlas-data').resolve()
    finally:
        os.chdir(previous)


def test_drive_search_scope_and_connect_requires_login(client):
    from pydantic import ValidationError

    from app.drive import SCOPE, DriveSearch

    assert SCOPE == 'https://www.googleapis.com/auth/drive.readonly'
    assert client.get('/api/drive/connect?wid=missing').status_code in (404, 503)
    with __import__('pytest').raises(ValidationError):
        DriveSearch(query='x' * 181)


def test_drive_search_builds_safe_query(monkeypatch):
    from app import drive

    captured = {}

    def fake_request(url, token, params=None, **kwargs):
        captured.update(params)
        return {'files': []}

    monkeypatch.setattr(drive, 'request_google', fake_request)
    assert drive.drive_search('test-token', "client's notes") == []
    assert "client\\'s notes" in captured['q']
    assert 'trashed = false' in captured['q']


def test_drive_evidence_keeps_title_when_file_has_no_extractable_text(monkeypatch):
    from app import drive

    monkeypatch.setattr(
        drive,
        'request_google',
        lambda *args, **kwargs: b'%PDF scanned image' if kwargs.get('binary') else '',
    )
    monkeypatch.setattr(drive, 'extract_drive_binary', lambda *_args: '')
    results = drive.drive_evidence(
        'token',
        [
            {
                'id': 'resume-1',
                'title': 'Resume 2026.pdf',
                'excerpt': 'Resume 2026.pdf',
                'mimeType': 'application/pdf',
            }
        ],
    )

    assert len(results) == 1
    assert results[0]['excerpt'] == 'Resume 2026.pdf'


def test_connected_chat_uses_standard_chat_citation_contract(client, monkeypatch):
    from app import connected_synthesis, main

    w = wid(client)

    def fake_retrieve(*args, **kwargs):
        return [
            {
                'chunk_id': 'chunk-123',
                'document_id': 'doc-123',
                'name': 'Plan.md',
                'page': 2,
                'location_type': 'section',
                'text': 'Agreed milestone',
            }
        ]

    def fake_synthesize(question, evidence, *args):
        return 'Milestone agreed [1].', [
            {**evidence[0].as_prompt(1), 'origin_id': evidence[0].source_id}
        ]

    monkeypatch.setattr(main, 'retrieve', fake_retrieve)
    monkeypatch.setattr(connected_synthesis, 'synthesize_connected', fake_synthesize)
    response = client.post(
        f'/api/workspaces/{w}/connected-chat', json={'question': 'What milestone was agreed?'}
    )
    assert response.status_code == 200, response.text
    result = response.json()
    assert result['mode'] == 'connected'
    assert result['sources'][0]['chunk_id'] == 'chunk-123'
    assert result['sources'][0]['document_id'] == 'doc-123'
    assert result['sources'][0]['page'] == 2
    assert result['sources'][0]['location_type'] == 'section'


def test_github_pull_request_listing_uses_connected_records_without_model(client, monkeypatch):
    from app import main
    from app.sources import Evidence

    w = wid(client)
    records = [
        Evidence(
            'github',
            'pr-1',
            'Fix account-scoped retrieval',
            'This PR scopes search to the authenticated account. More details follow.',
            'https://github.com/me/repo/pull/1',
        ),
        Evidence(
            'github',
            'pr-2',
            'Improve source citations',
            'Adds source links to each result.',
            'https://github.com/me/repo/pull/2',
        ),
    ]
    monkeypatch.setattr(main, 'collect_connected_evidence', lambda *args: (records, {}))
    monkeypatch.setattr(main, 'enabled', lambda: True)
    response = client.post(
        f'/api/workspaces/{w}/connected-chat',
        json={'question': 'my pull requests', 'include_documents': False, 'providers': ['github']},
    )
    assert response.status_code == 200, response.text
    result = response.json()
    assert result['mode'] == 'source_results'
    assert result['scope'] == ['github']
    assert 'Fix account-scoped retrieval' in result['answer']
    assert 'This PR scopes search to the authenticated account.' in result['answer']
    assert '[1]' in result['answer'] and '[2]' in result['answer']
    assert result['sources'][0]['url'] == records[0].url


def test_unqualified_github_search_is_scoped_to_connected_users_repositories(monkeypatch):
    from app import external_sources

    captured = {}

    def fake_http(url, token, provider, params=None, method='GET'):
        if url.endswith('/user'):
            return {'login': 'srbmaury'}
        captured.update(params or {})
        return {'items': []}

    monkeypatch.setattr(external_sources, 'http_json', fake_http)
    external_sources.source_search('github', 'token', 'What is deforestation?')
    assert captured['q'] == 'deforestation is:issue user:srbmaury'


def test_unqualified_github_pr_search_is_scoped_to_connected_users_repositories(monkeypatch):
    from app import external_sources

    captured = {}

    def fake_http(url, token, provider, params=None, method='GET'):
        if url.endswith('/user'):
            return {'login': 'srbmaury'}
        captured.update(params or {})
        return {'items': []}

    monkeypatch.setattr(external_sources, 'http_json', fake_http)
    external_sources.source_search('github', 'token', 'List open pull requests')
    assert captured['q'] == 'is:pr state:open user:srbmaury'


def test_health_and_processing_stats(client):
    workspace_id = wid(client)
    assert client.get('/api/health/live').json()['status'] == 'ok'
    assert client.get('/api/health/ready').json()['status'] == 'ready'
    response = client.get(f'/api/workspaces/{workspace_id}/processing-stats')
    assert response.status_code == 200
    assert response.json() == {'queued': 0, 'processing': 0, 'ready': 0, 'failed': 0}
    upload(client, workspace_id, 'metrics.md', 'A sufficiently long sample for metrics.')
    assert client.get(f'/api/workspaces/{workspace_id}/processing-stats').json()['queued'] == 1
    drain()


def test_retrieval_diversifies_sources():
    from app.retrieval import diversify_ranked_chunks

    ranked = [SimpleNamespace(id=f'a{i}', document_id='a') for i in range(6)]
    ranked += [SimpleNamespace(id=f'b{i}', document_id='b') for i in range(3)]
    picked = diversify_ranked_chunks(ranked, 6, per_document=3)
    assert len(picked) == 6
    assert len([x for x in picked if x.document_id == 'b']) == 3


def test_transient_worker_failure_requeues_with_delay(client, monkeypatch):
    from app import worker

    workspace_id = wid(client)
    document = upload(client, workspace_id, 'retry-transient.md', 'A sample document.')

    def unavailable(*args, **kwargs):
        raise RuntimeError('Temporary embedding outage')

    monkeypatch.setattr(worker, 'ingest', unavailable)
    assert worker.run_one()
    with Session() as db:
        state = db.get(Document, document['id'])
        assert state.status == 'queued'
        assert state.lease > time.time()
        assert state.generation == 1


def test_operations_endpoint_is_workspace_scoped(client):
    w = wid(client)
    response = client.get(f'/api/workspaces/{w}/operations')
    assert response.status_code == 200
    assert response.json()['documents']['total'] == 0
    upload(client, w, 'operations-note.txt', 'hello world')
    assert client.get(f'/api/workspaces/{w}/operations').json()['documents']['total'] == 1
    with TestClient(app) as guest:
        assert guest.get(f'/api/workspaces/{w}/operations').status_code == 401
        guest.post(
            '/api/auth/register',
            json={'email': f'{time.time_ns()}@guest.example', 'password': 'another-long-password'},
        )
        assert guest.get(f'/api/workspaces/{w}/operations').status_code == 404
